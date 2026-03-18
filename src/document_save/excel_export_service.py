import os
import logging
from datetime import datetime, date
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from src.queries.dto.execute_query_dto import ExecuteQueryDTO
from src.document_save.filename_service import FilenameService
from config import FileRepo

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


class ExcelExportService:
    base_path = FileRepo.base_drive

    def save_to_xlsx(self, results, query: ExecuteQueryDTO) -> str:
        """Generate .xlsx with Summary + Data sheets."""
        filename = FilenameService.generate_filename(
            user_id=query.user_id,
            query_name=query.name,
            query_params=query.query_params,
            timestamp=datetime.now(),
        ).replace(".csv", ".xlsx")

        file_path = os.path.join(
            self.base_path, "query_results",
            str(query.user_id), str(query.query_id), filename,
        )
        os.makedirs(os.path.dirname(file_path), exist_ok=True)

        wb = Workbook()

        # --- Summary Sheet ---
        ws_summary = wb.active
        ws_summary.title = "Summary"
        ws_summary["A1"] = "NIB Query Tool — Report Summary"
        ws_summary["A1"].font = Font(bold=True, size=14)
        ws_summary["A3"] = "Query Name:"
        ws_summary["B3"] = query.name
        ws_summary["A4"] = "Generated:"
        ws_summary["B4"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws_summary["A5"] = "Total Rows:"
        ws_summary["B5"] = len(results.rows) if results and results.rows else 0

        # Numeric column summaries
        row_offset = 7
        if results and results.rows and results.column_names:
            for col_idx, col_name in enumerate(results.column_names):
                sample = [r[col_idx] for r in results.rows[:100] if r[col_idx] is not None]
                if sample and all(isinstance(v, (int, float, Decimal)) for v in sample):
                    values = [float(v) for v in [r[col_idx] for r in results.rows] if v is not None]
                    if values:
                        ws_summary[f"A{row_offset}"] = f"{col_name} (Total):"
                        ws_summary[f"B{row_offset}"] = sum(values)
                        ws_summary[f"A{row_offset + 1}"] = f"{col_name} (Avg):"
                        ws_summary[f"B{row_offset + 1}"] = round(sum(values) / len(values), 2)
                        row_offset += 2

        ws_summary.column_dimensions["A"].width = 30
        ws_summary.column_dimensions["B"].width = 40

        # --- Data Sheet ---
        ws_data = wb.create_sheet("Data")

        # Headers
        for col_idx, col_name in enumerate(results.column_names, 1):
            cell = ws_data.cell(row=1, column=col_idx, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        # Data rows
        for row_idx, row in enumerate(results.rows, 2):
            for col_idx, value in enumerate(row, 1):
                cell = ws_data.cell(row=row_idx, column=col_idx)
                if isinstance(value, Decimal):
                    cell.value = float(value)
                elif isinstance(value, (date, datetime)):
                    cell.value = value
                    cell.number_format = "YYYY-MM-DD"
                else:
                    cell.value = value
                cell.border = THIN_BORDER

        # Auto-width columns
        for col_idx in range(1, len(results.column_names) + 1):
            max_len = len(str(results.column_names[col_idx - 1]))
            for row in results.rows[:50]:
                val_len = len(str(row[col_idx - 1])) if row[col_idx - 1] else 0
                max_len = max(max_len, val_len)
            ws_data.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

        # Freeze header row
        ws_data.freeze_panes = "A2"

        try:
            wb.save(file_path)
            logger.info("Excel report saved: %s (%d rows)", file_path, len(results.rows))
        except Exception as e:
            logger.error("Error saving Excel: %s", e, exc_info=True)
            raise

        return file_path
